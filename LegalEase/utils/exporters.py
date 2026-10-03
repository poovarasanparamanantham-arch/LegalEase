import io
import re
from pathlib import Path
from typing import Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from fpdf import FPDF


def text_to_txt(text: str) -> bytes:
    return text.encode("utf-8")


def _clean_lines(text: str):
    return [line.strip() for line in text.replace("\r\n", "\n").split("\n")]


def format_docx(text: str, doc_type: str, logo_path: Optional[str] = None) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)

    if logo_path and Path(logo_path).exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(logo_path, width=Inches(1.25))

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(doc_type.upper())
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(15)

    for line in _clean_lines(text):
        if not line:
            doc.add_paragraph()
            continue

        if re.match(r"^(\d+[\.\)]|SECTION\s+\d+|ARTICLE\s+\d+)", line, re.I):
            p = doc.add_paragraph()
            r = p.add_run(line)
            r.bold = True
            r.font.name = "Times New Roman"
            r.font.size = Pt(11)
        else:
            p = doc.add_paragraph(line)
            p.paragraph_format.space_after = Pt(6)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Generated with LegalEase - AI drafting aid; review with a qualified legal professional.")

    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


class LegalPDF(FPDF):
    def __init__(self, logo_path: Optional[str] = None):
        super().__init__()
        self.logo_path = logo_path

    def header(self):
        if self.logo_path and Path(self.logo_path).exists():
            try:
                self.image(self.logo_path, x=92, y=8, w=26)
                self.ln(22)
            except Exception:
                self.ln(5)
        else:
            self.ln(5)

    def footer(self):
        self.set_y(-15)
        self.set_font("Helvetica", size=8)
        self.cell(
            0,
            8,
            "LegalEase - AI drafting aid; review with a qualified legal professional.",
            align="C",
        )


def format_pdf(text: str, doc_type: str, logo_path: Optional[str] = None) -> bytes:
    pdf = LegalPDF(logo_path=logo_path)
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.add_page()
    pdf.set_title(doc_type)

    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 12, doc_type.upper(), align="C")
    pdf.ln(14)

    for line in _clean_lines(text):
        if not line:
            pdf.ln(3)
            continue

        is_heading = bool(re.match(r"^(\d+[\.\)]|SECTION\s+\d+|ARTICLE\s+\d+)", line, re.I))
        pdf.set_font("Helvetica", "B" if is_heading else "", 11)
        safe = line.encode("latin-1", "replace").decode("latin-1")
        pdf.multi_cell(0, 6, safe)
        pdf.ln(1)

    return bytes(pdf.output(dest="S"))


def format_html_preview(text: str) -> str:
    import html

    escaped = html.escape(text)
    blocks = []
    for paragraph in escaped.split("\n\n"):
        if paragraph.strip():
            blocks.append(f"<p>{paragraph.replace(chr(10), '<br>')}</p>")
    return "\n".join(blocks)
